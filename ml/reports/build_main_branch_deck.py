"""Generates ml/reports/wifi_csi_main_branch_model.pptx -- summarizes the model actually committed to
main: WhoFiTransformer (ml/models/transformer_whofi.py), trained per-task by
ml/training/train_final_model.py / train_day3_ch6_model.py, served by ml/inference/live_infer.py.

This is NOT the RandomForest "home model" or the Signature+EVM open-set model -- those only exist as
uncommitted work-in-progress files on top of this branch. This deck covers what's actually merged:
a single-branch (amplitude-only) transformer classifier, trained independently for 3-4 binary tasks
(presence / authorized-vs-not / motion / legacy identity), with per-day empty-room (calibA) calibration.

Source: ml/models/transformer_whofi.py, ml/models/common.py, ml/data_pipeline/{decode_csi,csi_resample,
windowing,calibration}.py, ml/training/{train,train_final_model,train_day3_ch6_model}.py,
ml/inference/{live_infer,live_window,live_calibration,checkpoint}.py, ml/reports/day2_next_steps.md.

Run: .venv-ml/Scripts/python.exe ml/reports/build_main_branch_deck.py
"""
from __future__ import annotations

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.oxml.ns import qn

# ---- palette (validated categorical set, light-surface steps) --------------------------------
INK = RGBColor(0x0B, 0x0B, 0x0B)
INK_SECONDARY = RGBColor(0x52, 0x51, 0x4E)
MUTED = RGBColor(0x89, 0x87, 0x81)
SURFACE = RGBColor(0xFC, 0xFC, 0xFB)
PAGE = RGBColor(0xF9, 0xF9, 0xF7)
BORDER = RGBColor(0xC3, 0xC2, 0xB7)

BLUE = RGBColor(0x2A, 0x78, 0xD6)
ORANGE = RGBColor(0xEB, 0x68, 0x34)
AQUA = RGBColor(0x1B, 0xAF, 0x7A)
YELLOW = RGBColor(0xED, 0xA1, 0x00)
VIOLET = RGBColor(0x4A, 0x3A, 0xA7)
RED = RGBColor(0xE3, 0x49, 0x48)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
NAVY = RGBColor(0x11, 0x1E, 0x2E)

SLIDE_W, SLIDE_H = Inches(13.333), Inches(7.5)

prs = Presentation()
prs.slide_width = SLIDE_W
prs.slide_height = SLIDE_H
BLANK = prs.slide_layouts[6]


def add_slide():
    slide = prs.slides.add_slide(BLANK)
    bg = slide.background
    bg.fill.solid()
    bg.fill.fore_color.rgb = PAGE
    return slide


def add_rect(slide, x, y, w, h, fill=None, line=None, line_w=Pt(1)):
    shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
    try:
        shp.adjustments[0] = 0.06
    except Exception:
        pass
    if fill is None:
        shp.fill.background()
    else:
        shp.fill.solid()
        shp.fill.fore_color.rgb = fill
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line
        shp.line.width = line_w
    shp.shadow.inherit = False
    return shp


def set_text(shape, text, size=14, color=INK, bold=False, align=PP_ALIGN.LEFT,
             anchor=MSO_ANCHOR.MIDDLE, font="Segoe UI", wrap=True):
    tf = shape.text_frame
    tf.word_wrap = wrap
    tf.vertical_anchor = anchor
    tf.margin_left = Pt(6)
    tf.margin_right = Pt(6)
    tf.margin_top = Pt(4)
    tf.margin_bottom = Pt(4)
    lines = text.split("\n")
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        run = p.add_run()
        run.text = line
        run.font.size = Pt(size)
        run.font.color.rgb = color
        run.font.bold = bold
        run.font.name = font
    return shape


def add_textbox(slide, x, y, w, h, text, size=14, color=INK, bold=False,
                 align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(x, y, w, h)
    return set_text(box, text, size=size, color=color, bold=bold, align=align, anchor=anchor)


def add_bullets(slide, x, y, w, h, items, size=15, color=INK, line_spacing=1.15):
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        if isinstance(item, tuple):
            text, level = item
        else:
            text, level = item, 0
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.level = level
        p.line_spacing = line_spacing
        p.space_after = Pt(6)
        marker = "•  " if level == 0 else "–  "
        run = p.add_run()
        run.text = marker + text
        run.font.size = Pt(size - level * 1)
        run.font.color.rgb = color if level == 0 else INK_SECONDARY
        run.font.name = "Segoe UI"
    return box


def add_kicker_title(slide, kicker, title, accent=BLUE):
    add_rect(slide, Inches(0.55), Inches(0.35), Inches(0.09), Inches(0.85), fill=accent)
    add_textbox(slide, Inches(0.8), Inches(0.32), Inches(9), Inches(0.35), kicker.upper(),
                size=12, color=accent, bold=True)
    add_textbox(slide, Inches(0.8), Inches(0.6), Inches(11.7), Inches(0.65), title,
                size=26, color=INK, bold=True)


def add_footer(slide, n):
    add_textbox(slide, Inches(0.55), Inches(7.12), Inches(7), Inches(0.3),
                "WiFi-CSI Presence Collector — Main-Branch Model (WhoFi Transformer)",
                size=9, color=MUTED)
    add_textbox(slide, Inches(12.3), Inches(7.12), Inches(0.5), Inches(0.3), str(n),
                size=9, color=MUTED, align=PP_ALIGN.RIGHT)


def add_arrow(slide, x1, y1, x2, y2, color=MUTED, width=Pt(1.5)):
    conn = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, x1, y1, x2, y2)
    conn.line.color.rgb = color
    conn.line.width = width
    line = conn.line._get_or_add_ln()
    arrow = line.makeelement(qn("a:tailEnd"), {"type": "triangle", "w": "med", "len": "med"})
    line.append(arrow)
    return conn


def flow_step(slide, x, y, w, h, label, sub=None, fill=SURFACE, line=BORDER, label_color=INK,
              label_size=13, sub_size=10):
    box = add_rect(slide, x, y, w, h, fill=fill, line=line, line_w=Pt(1.25))
    if sub:
        tf = box.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        for i, line in enumerate(label.split("\n")):
            p0 = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p0.alignment = PP_ALIGN.CENTER
            r0 = p0.add_run()
            r0.text = line
            r0.font.size = Pt(label_size)
            r0.font.bold = True
            r0.font.color.rgb = label_color
            r0.font.name = "Segoe UI"
        for line in sub.split("\n"):
            p1 = tf.add_paragraph()
            p1.alignment = PP_ALIGN.CENTER
            r1 = p1.add_run()
            r1.text = line
            r1.font.size = Pt(sub_size)
            r1.font.color.rgb = INK_SECONDARY
            r1.font.name = "Segoe UI"
    else:
        set_text(box, label, size=label_size, color=label_color, bold=True, align=PP_ALIGN.CENTER)
    return box


# =================================================================================================
# Slide 1 -- Title
# =================================================================================================
s = add_slide()
add_rect(s, 0, 0, SLIDE_W, Inches(7.5), fill=NAVY)
add_rect(s, Inches(0.9), Inches(2.5), Inches(0.12), Inches(1.9), fill=BLUE)
add_textbox(s, Inches(1.2), Inches(2.45), Inches(11), Inches(0.4), "WIFI-CSI PRESENCE COLLECTOR — MAIN BRANCH",
            size=14, color=RGBColor(0x9E, 0xC5, 0xF4), bold=True)
add_textbox(s, Inches(1.15), Inches(2.85), Inches(11.5), Inches(1.6),
            "The Actual Shipped Model:\nWhoFi Transformer + Empty-Room Calibration",
            size=34, color=WHITE, bold=True)
add_textbox(s, Inches(1.2), Inches(4.55), Inches(10.5), Inches(0.9),
            "Not the Random Forest “home model” and not the Signature+EVM open-set experiment — both of\n"
            "those are later, uncommitted work. This is what's actually merged into main.",
            size=15, color=RGBColor(0xC3, 0xC2, 0xB7))
add_textbox(s, Inches(1.2), Inches(6.6), Inches(10), Inches(0.4),
            "ml/models/transformer_whofi.py — train_final_model.py / train_day3_ch6_model.py — live_infer.py",
            size=11, color=MUTED)

add_footer(s, 1)

# =================================================================================================
# Slide 2 -- Overview
# =================================================================================================
s = add_slide()
add_kicker_title(s, "Overview", "What's actually on main: one transformer, four tasks", accent=BLUE)

add_bullets(s, Inches(0.8), Inches(1.35), Inches(11.7), Inches(5.4), [
    "Model: WhoFiTransformer — a replica of WhoFi (arXiv:2507.12869): 1-layer transformer over "
    "amplitude only, mean-pooled, L2-normalized, then a linear classifier head",
    "Trained independently as a plain closed-set classifier for each of 4 binary tasks "
    "(cross-entropy loss, not contrastive) — there is no open-set / EVM layer on main",
    ("task0_presence: EMPTY vs OCCUPIED", 1),
    ("taskD_auth_vs_nonauth: NOT AUTHORIZED vs AUTHORIZED  (the “auth gate”)", 1),
    ("taskE_motion_standing_vs_walking: STANDING vs WALKING", 1),
    ("taskB_identity (legacy, ANJALI vs BARATH): kept because it was asked for, but tested BELOW CHANCE "
     "cross-day (AUROC 0.20–0.31) — live_infer.py marks it experimental/unreliable", 1),
    "Preprocessing: per-day empty-room self-calibration (“calibA”) — z-score every window against "
    "that day's own quiet-room baseline before the transformer ever sees it",
    "Trained on ALL of Day1+Day2 pooled data, no held-out split — Day3 live testing is the real test "
    "(train_day3_ch6_model.py later added its own Day3+Day4-only checkpoint track)",
    "Served by ml/inference/live_infer.py: live calibration → rolling 200-packet windows → "
    "per-task prediction → ~30s rolling aggregate, which is the number to trust",
], size=15)

add_footer(s, 2)

# =================================================================================================
# Slide 3 -- End-to-end flowchart
# =================================================================================================
s = add_slide()
add_kicker_title(s, "Flowchart", "Raw packets → preprocessing → transformer → live decision", accent=BLUE)

steps_row1 = [
    ("Raw CSI packet", "ragged csi_len, per-802.11-frame-type", NAVY),
    ("Decode", "amp=hypot(re,im)\nphase=atan2(im,re)", AQUA),
    ("Subcarrier resample", "Day1 186-sub, Day2 128-sub\n→ shared 128-sub grid", AQUA),
]
sw, sgap = Inches(3.55), Inches(0.25)
stotal = sw * 3 + sgap * 2
sx0 = int((SLIDE_W - stotal) / 2)
y1 = Inches(1.25)
for i, (label, sub, color) in enumerate(steps_row1):
    x = sx0 + i * (sw + sgap)
    flow_step(s, x, y1, sw, Inches(0.75), label, sub, fill=color, line=color, label_color=WHITE)
    if i < 2:
        add_arrow(s, x + sw, y1 + Inches(0.375), x + sw + sgap, y1 + Inches(0.375))

add_arrow(s, SLIDE_W / 2, y1 + Inches(0.75), SLIDE_W / 2, y1 + Inches(0.95))

steps_row2 = [
    ("calibA calibration", "z = (x − μ_empty_day) / σ_empty_day\namplitude + phase", ORANGE),
    ("Windowing", "200-packet window,\n100-packet stride (50% overlap)", ORANGE),
    ("WhoFiTransformer", "1-layer self-attention\n→ 32-d embedding", BLUE),
]
y2 = y1 + Inches(0.95) + Inches(0.2)
for i, (label, sub, color) in enumerate(steps_row2):
    x = sx0 + i * (sw + sgap)
    flow_step(s, x, y2, sw, Inches(0.85), label, sub, fill=color, line=color, label_color=WHITE)
    if i < 2:
        add_arrow(s, x + sw, y2 + Inches(0.425), x + sw + sgap, y2 + Inches(0.425))

add_arrow(s, SLIDE_W / 2, y2 + Inches(0.85), SLIDE_W / 2, y2 + Inches(1.05))

# 4 task heads
y3 = y2 + Inches(1.05) + Inches(0.2)
add_textbox(s, sx0, y3 - Inches(0.05), stotal, Inches(0.3),
            "4 independent linear classifier heads (one per task, same shared architecture)",
            size=12.5, color=INK_SECONDARY, align=PP_ALIGN.CENTER)
tasks = [
    ("Presence", "EMPTY / OCCUPIED"),
    ("Auth gate", "AUTH / NOT AUTH"),
    ("Motion", "STAND / WALK"),
    ("Identity*", "ANJALI / BARATH"),
]
tw = Inches(2.75)
tgap = Inches(0.2)
ttotal = tw * 4 + tgap * 3
tx0 = int((SLIDE_W - ttotal) / 2)
ty = y3 + Inches(0.3)
for i, (label, sub) in enumerate(tasks):
    x = tx0 + i * (tw + tgap)
    color = MUTED if i == 3 else AQUA
    flow_step(s, x, ty, tw, Inches(0.7), label, sub, fill=SURFACE, line=color, label_color=INK)

add_arrow(s, SLIDE_W / 2, ty + Inches(0.7), SLIDE_W / 2, ty + Inches(0.9))
flow_step(s, sx0, ty + Inches(0.95), stotal, Inches(0.55),
          "Live: ~30s rolling aggregate across overlapping windows — the number to trust",
          fill=NAVY, line=NAVY, label_color=WHITE, label_size=13)

add_textbox(s, sx0, ty + Inches(1.65), stotal, Inches(0.3),
            "* taskB identity tested below chance cross-day — shown for completeness, not trusted",
            size=10.5, color=MUTED, align=PP_ALIGN.CENTER)

add_footer(s, 3)

# =================================================================================================
# Slide 4 -- Raw CSI decode
# =================================================================================================
s = add_slide()
add_kicker_title(s, "Raw data", "Decoding the ragged CSI wire format", accent=BLUE)

add_bullets(s, Inches(0.8), Inches(1.4), Inches(5.7), Inches(4.9), [
    "samples.npz stores CSI as a flat, ragged buffer — csi_len varies by 802.11 frame type "
    "(legacy/HT, 20/40 MHz), so packets are bucketed by csi_len before decoding",
    "Each subcarrier is one (imag, real) int8 pair; decode is vectorized per bucket (gather + reshape), "
    "not a per-packet Python loop",
    "amplitude = hypot(real, imag) = sqrt(real² + imag²)",
    "phase = arctan2(imag, real)",
    "Day 1 sessions negotiated 40 MHz → dominant bucket is 186 subcarriers (372 bytes)",
    "Day 2 (and Day 3) sessions negotiated 20 MHz → dominant bucket is 128 subcarriers (256 bytes)",
    "Only the per-session dominant csi_len bucket is kept (~1–6% of minority-shape packets dropped) "
    "so every window has one consistent subcarrier count",
], size=14)

diag_x = Inches(7.0)
add_rect(s, diag_x, Inches(1.4), Inches(5.4), Inches(4.9), fill=SURFACE, line=BORDER)
add_textbox(s, diag_x + Inches(0.3), Inches(1.55), Inches(4.8), Inches(0.35),
            "Shape mismatch across days", size=13.5, color=INK, bold=True)
flow_step(s, diag_x + Inches(0.4), Inches(2.05), Inches(2.2), Inches(0.85),
          "Day 1", "40 MHz\n186 subcarriers", fill=VIOLET, line=VIOLET, label_color=WHITE)
flow_step(s, diag_x + Inches(2.8), Inches(2.05), Inches(2.2), Inches(0.85),
          "Day 2 / Day 3", "20 MHz\n128 subcarriers", fill=AQUA, line=AQUA, label_color=WHITE)
add_arrow(s, diag_x + Inches(1.5), Inches(2.9), diag_x + Inches(2.7), Inches(3.55))
add_arrow(s, diag_x + Inches(3.9), Inches(2.9), diag_x + Inches(2.9), Inches(3.55))
flow_step(s, diag_x + Inches(0.7), Inches(3.6), Inches(4.0), Inches(0.75),
          "resample onto shared 128-sub grid", "downsample Day1 only — never fabricate Day2's missing resolution",
          fill=BLUE, line=BLUE, label_color=WHITE)
add_textbox(s, diag_x + Inches(0.3), Inches(4.7), Inches(4.8), Inches(1.4),
            "Why 128, not 186: resampling always shrinks toward the smaller native width so both days "
            "become shape-compatible for one shared model — upsampling Day 2 would invent detail the "
            "hardware never captured.",
            size=12.5, color=INK_SECONDARY)

add_footer(s, 4)

# =================================================================================================
# Slide 5 -- calibA calibration
# =================================================================================================
s = add_slide()
add_kicker_title(s, "Preprocessing", "calibA — per-day empty-room self-calibration", accent=ORANGE)

add_rect(s, Inches(0.8), Inches(1.3), Inches(11.7), Inches(1.35), fill=SURFACE, line=BORDER)
add_textbox(s, Inches(1.1), Inches(1.42), Inches(11), Inches(0.35),
            "Per subcarrier, applied to both amplitude and phase:", size=13.5, color=INK, bold=True)
add_textbox(s, Inches(1.1), Inches(1.8), Inches(11), Inches(0.7),
            "x' = ( x − μ_empty ) / ( σ_empty + ε )",
            size=22, color=BLUE, bold=True)

add_bullets(s, Inches(0.8), Inches(2.95), Inches(11.7), Inches(3.9), [
    "μ_empty, σ_empty are fit once per day from that day's own “none” (empty-room) sessions — "
    "amplitude and phase each get their own per-subcarrier mean/std",
    "Why per-day, not global: this is the direct analog of OpenCSI's Z-score approach (arXiv:2607.26665), "
    "which reported F1 0.99 vs 0.87 for cross-room/cross-hardware transfer on ESP32 — tested here for "
    "person-ID/auth for the first time",
    "Live inference computes the same baseline on the fly: a 15s “stand outside the room” calibration "
    "phase at the start of every session (ml/inference/live_calibration.py), reusing the exact same "
    "DayBaseline math and apply_variant_a() function as the offline training pipeline",
    "Two alternatives were tried and NOT used as the default:",
    ("Variant B (cross-day re-referencing onto a fixed reference day) — underperforms calibA, "
     "likely because it assumes the empty-vs-occupied relationship transforms identically across days", 1),
    ("ℓ₁ per-packet gain normalization — implemented and unit-tested, not yet run through the full sweep", 1),
    "Net effect measured (see results slide): calibA trades some raw AUROC for a much lower "
    "false-accept rate against a genuine stranger — the metric that actually matters for the auth task",
], size=14)

add_footer(s, 5)

# =================================================================================================
# Slide 6 -- Windowing
# =================================================================================================
s = add_slide()
add_kicker_title(s, "Preprocessing", "Windowing — turning a packet stream into fixed-size inputs", accent=ORANGE)

steps = [
    ("Calibrated stream", "(N packets, 128 subcarriers)\namplitude + phase"),
    ("Slide a window", "window_packets = 200\n(~1s at ~200Hz observed rate)"),
    ("Stride forward", "stride_packets = 100\n(50% overlap)"),
    ("Model input", "(200, 128)\nper window"),
]
step_w = Inches(2.75)
gap = Inches(0.3)
total_w = step_w * len(steps) + gap * (len(steps) - 1)
start_x = int((SLIDE_W - total_w) / 2)
y = Inches(1.7)
for i, (label, sub) in enumerate(steps):
    x = start_x + i * (step_w + gap)
    flow_step(s, x, y, step_w, Inches(1.3), label, sub, fill=SURFACE, line=ORANGE)
    if i < len(steps) - 1:
        add_arrow(s, x + step_w, y + Inches(0.65), x + step_w + gap, y + Inches(0.65))

add_bullets(s, Inches(0.8), Inches(3.5), Inches(11.7), Inches(3.1), [
    "Observed packet rate on the dominant bucket: 162–263 Hz (mean ~217 Hz) across Day1 sessions — "
    "so 200 packets is close to a 1-second window, matching the ESP32 person-ID paper's convention",
    "The live rolling windower (ml/inference/live_window.py) mirrors this exactly, so a live window is "
    "the same shape a checkpoint was trained on — the very first window is ready as soon as the "
    "buffer first fills, not after one extra stride's delay",
    "Per-session caching avoids re-decoding: each session's dominant-bucket arrays are decoded once "
    "and memory-mapped, so building thousands of overlapping windows stays cheap",
    "train_day3_ch6_model.py's later checkpoint track adds one more step before windowing: bin-average "
    "resampling of the TIME axis onto a common packet rate, to stop window duration from correlating "
    "with whichever label happened to be recorded at a different network congestion level",
], size=14)

add_footer(s, 6)

# =================================================================================================
# Slide 7 -- Model architecture
# =================================================================================================
s = add_slide()
add_kicker_title(s, "Model architecture", "WhoFiTransformer — single amplitude branch + linear head", accent=BLUE)

diag_x = Inches(0.8)
diag_w = Inches(5.6)
y = Inches(1.25)
box_h = Inches(0.62)
gap = Inches(0.14)

layers = [
    ("Input: amplitude (B, T=200, S=128)", BLUE),
    ("BranchEncoder: Linear 128→32 + positional encoding", AQUA),
    ("Transformer encoder × 1\n(self-attn → Add&Norm → FFN → Add&Norm)", ORANGE),
    ("Mean-pool over time → (B, 32)", VIOLET),
    ("signature_module: Linear 32→32, L2-normalize", VIOLET),
    ("classifier: Linear 32→n_classes  (per task)", BLUE),
    ("Output: logits (B, n_classes)", NAVY),
]
cy = y
for i, (label, color) in enumerate(layers):
    h = Inches(0.85) if "\n" in label else box_h
    box = flow_step(s, diag_x, cy, diag_w, h, label, fill=color, line=color, label_color=WHITE,
                     label_size=12)
    cy = cy + h
    if i < len(layers) - 1:
        add_arrow(s, diag_x + diag_w / 2, cy, diag_x + diag_w / 2, cy + gap)
        cy = cy + gap

rx = Inches(6.75)
add_rect(s, rx, Inches(1.25), Inches(5.75), Inches(5.6), fill=SURFACE, line=BORDER)
add_textbox(s, rx + Inches(0.3), Inches(1.4), Inches(5.15), Inches(0.4),
            "Notes on this exact architecture", size=14.5, color=INK, bold=True)
add_bullets(s, rx + Inches(0.3), Inches(1.85), Inches(5.15), Inches(4.9), [
    "Amplitude only, no phase branch — matches WhoFi's own published design; phase is accepted by "
    "forward() for a uniform call signature across the model zoo but is otherwise ignored",
    "d_model=32, 4 heads, d_ff=64, 1 layer — WhoFi's own ablation found 3 layers unstable/worse on "
    "a dataset this small, so the shallow config is the deliberate winner, not a placeholder",
    "signature_module (Linear 32→32 + L2-normalize) produces an embedding shaped like a "
    "verification “signature”, but on main it feeds straight into a plain linear classifier — "
    "there's no contrastive loss or EVM open-set layer wired up here (that's the later, uncommitted "
    "signature+EVM experiment)",
    "One separate model instance is trained per task (presence / auth / motion / identity) — same "
    "architecture and hyperparameters, four independent checkpoints, not a shared multi-head model",
    "classifier is a single nn.Linear(32, n_classes) — as simple as the architecture gets after the "
    "transformer; all the representational work happens in the 1-layer encoder",
], size=13)

add_footer(s, 7)

# =================================================================================================
# Slide 8 -- Self-attention math
# =================================================================================================
s = add_slide()
add_kicker_title(s, "Transformer internals", "Scaled dot-product multi-head self-attention", accent=BLUE)

add_rect(s, Inches(0.8), Inches(1.3), Inches(11.7), Inches(1.9), fill=SURFACE, line=BORDER)
add_textbox(s, Inches(1.1), Inches(1.45), Inches(11), Inches(0.35), "Per head:", size=14, color=INK, bold=True)
add_textbox(s, Inches(1.1), Inches(1.85), Inches(11), Inches(0.55),
            "Attention(Q, K, V) = softmax( Q Kᵀ / √d_k ) V", size=20, color=BLUE, bold=True)
add_textbox(s, Inches(1.1), Inches(2.5), Inches(11), Inches(0.6),
            "Q, K, V are learned linear projections of the 32-dim per-packet token sequence; "
            "d_k = d_model / n_heads = 32 / 4 = 8 here.", size=13, color=INK_SECONDARY)

add_bullets(s, Inches(0.8), Inches(3.4), Inches(5.7), Inches(3.5), [
    "Each of the 200 packets in a window becomes one 32-dim token via the input projection",
    "4 heads run attention in parallel, then concatenate — each head can latch onto different "
    "subcarrier/time relationships across the window",
    "Post-LN block structure (this repo's default):",
    ("x = LayerNorm( x + Dropout( SelfAttention(x) ) )", 1),
    ("x = LayerNorm( x + Dropout( FFN(x) ) )", 1),
], size=14)

add_bullets(s, Inches(6.85), Inches(3.4), Inches(5.65), Inches(3.5), [
    "Sinusoidal positional encoding, added elementwise to the projected input:",
    ("PE(pos, 2i)   = sin( pos / 10000^(2i/d_model) )", 1),
    ("PE(pos, 2i+1) = cos( pos / 10000^(2i/d_model) )", 1),
    "Mean-pool collapses the 200 refined tokens to one (B, 32) vector per window",
    "L2-normalize (signature_module output): s = z / ‖z‖₂ → ‖s‖ = 1 — "
    "used here only as an internal representation before the linear classifier, not for distance-based "
    "matching (no EVM on main)",
], size=14)

add_footer(s, 8)

# =================================================================================================
# Slide 9 -- Training regime
# =================================================================================================
s = add_slide()
add_kicker_title(s, "Training", "Four independent binary classifiers, one recipe", accent=ORANGE)

add_bullets(s, Inches(0.8), Inches(1.35), Inches(11.7), Inches(5.5), [
    "Loss: plain torch.nn.CrossEntropyLoss — closed-set classification, not contrastive/in-batch-negative "
    "(that loss exists in ml/training/losses.py but is not what trains the shipped checkpoints)",
    "Optimizer: Adam, lr=1e-3, batch_size=64, seed fixed and logged per run",
    ("Why the seed matters: re-training the SAME model on the SAME fold gave wildly different results "
     "(taskD fold 1: AUROC 0.757 vs 0.698, false-accept-unauthorized 26% vs 57%, across two otherwise-"
     "identical runs) — on a dataset this small with only a handful of epochs, seed variance can be as "
     "large as the effect being measured", 1),
    "Deployable checkpoints (train_final_model.py): trained on ALL Day1+Day2 pooled data, the SAME "
    "dataset passed as both train and test — deliberate, not an oversight",
    ("The printed accuracy is an in-sample sanity check only, never a generalization estimate — "
     "carving out an internal validation slice would reintroduce the 50%-overlapping-window leakage "
     "problem session-disjoint splitting exists to prevent", 1),
    ("Day 3, tested live via live_infer.py, is the real held-out test", 1),
    "Later track (train_day3_ch6_model.py): honest numbers come FIRST — leave-one-unauthorized-person-out "
    "for the auth task, session-disjoint 80/20 for presence/motion — THEN a final checkpoint is retrained "
    "on 100% of that track's data, so the two-stage “evaluate honestly, then ship” discipline is "
    "consistent across both checkpoint generations",
], size=14.5)

add_footer(s, 9)

# =================================================================================================
# Slide 10 -- Live inference pipeline
# =================================================================================================
s = add_slide()
add_kicker_title(s, "Live inference", "live_infer.py — from board stream to a trusted decision", accent=ORANGE)

steps = [
    ("Bandwidth/channel gate", "checked BEFORE the calibration\nwait — fail fast on wrong config"),
    ("Live calibration", "~15s “stand outside\nthe room” → empty-room baseline"),
    ("Rolling window", "200-packet window,\n100-packet stride"),
    ("4 task predictions", "instantaneous\nprobability per task"),
    ("~30s rolling aggregate", "mean over ~60 windows —\nthe number to trust"),
]
step_w = Inches(2.25)
gap = Inches(0.18)
total_w = step_w * len(steps) + gap * (len(steps) - 1)
start_x = int((SLIDE_W - total_w) / 2)
y = Inches(1.5)
for i, (label, sub) in enumerate(steps):
    x = start_x + i * (step_w + gap)
    flow_step(s, x, y, step_w, Inches(1.3), label, sub, fill=SURFACE, line=ORANGE, label_size=12, sub_size=9.5)
    if i < len(steps) - 1:
        add_arrow(s, x + step_w, y + Inches(0.65), x + step_w + gap, y + Inches(0.65))

add_bullets(s, Inches(0.8), Inches(3.25), Inches(11.7), Inches(3.6), [
    "Presence gates the display of Auth/Motion: if the aggregate presence probability says “empty”, "
    "Auth and Motion print “n/a (no one present)” rather than a misleading guess",
    "Guided phases (optional): prompts + a live countdown for “STAND STILL” then “WALK AROUND”, so a "
    "tester knows exactly when to switch — Auth doesn't need this split, Motion does",
    "Time-normalized checkpoints (the day3ch6 track) additionally bin-average the raw live packet stream "
    "onto the training-time target rate before windowing, via LiveTimeResampler — otherwise the exact "
    "packet-rate confound the offline fix closed would reappear live",
    "A periodic bandwidth re-check runs every 3000 packets during streaming — warns rather than aborts, "
    "since aborting a live demo mid-flight over a possibly-transient blip is worse than a warning",
], size=14)

add_footer(s, 10)

# =================================================================================================
# Slide 11 -- Validated results
# =================================================================================================
s = add_slide()
add_kicker_title(s, "Validated results", "raw vs. calibA, and why aggregation window length matters", accent=RED)

rows = [
    ("Preprocessing", "n_windows", "Accuracy", "AUROC", "False-accept unauth.", "False-accept none"),
    ("raw", "60 (~30s)", "0.798", "0.899", "58.2%", "3.9%"),
    ("calibA", "60 (~30s)", "0.749", "0.821", "5.5%", "1.6%"),
]
table_x, table_y = Inches(0.8), Inches(1.35)
table_w, table_h = Inches(11.7), Inches(1.7)
gtable = s.shapes.add_table(len(rows), len(rows[0]), table_x, table_y, table_w, table_h).table
widths = [Inches(2.3), Inches(1.7), Inches(1.7), Inches(1.6), Inches(2.2), Inches(2.2)]
for i, w in enumerate(widths):
    gtable.columns[i].width = w
for r, row in enumerate(rows):
    for c, val in enumerate(row):
        cell = gtable.cell(r, c)
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        cell.margin_left = Pt(6)
        cell.margin_right = Pt(6)
        para = cell.text_frame.paragraphs[0]
        para.alignment = PP_ALIGN.CENTER if c > 0 else PP_ALIGN.LEFT
        run = para.add_run()
        run.text = val
        run.font.size = Pt(13 if r > 0 else 13)
        run.font.name = "Segoe UI"
        if r == 0:
            run.font.bold = True
            run.font.color.rgb = WHITE
            cell.fill.solid()
            cell.fill.fore_color.rgb = NAVY
        else:
            run.font.bold = (c == 0) or (r == 2 and c == 4)
            run.font.color.rgb = RED if (r == 2 and c == 4) else INK_SECONDARY
            cell.fill.solid()
            cell.fill.fore_color.rgb = SURFACE if r % 2 else PAGE

add_textbox(s, Inches(0.8), Inches(3.25), Inches(11.7), Inches(0.4),
            "Calibration and aggregation window length are COMPOUNDING levers — false-accept-unauthorized "
            "under calibA shrinks monotonically as the window lengthens:", size=13.5, color=INK, bold=True)

fa_steps = [("1 window\n(~1s)", "21.5%"), ("~5s", "13.7%"), ("~10s", "13.0%"), ("~20s", "8.0%"), ("~30s", "5.5%")]
fw = Inches(2.15)
fgap = Inches(0.2)
ftotal = fw * len(fa_steps) + fgap * (len(fa_steps) - 1)
fx0 = int((SLIDE_W - ftotal) / 2)
fy = Inches(3.75)
for i, (label, pct) in enumerate(fa_steps):
    x = fx0 + i * (fw + fgap)
    is_last = i == len(fa_steps) - 1
    flow_step(s, x, fy, fw, Inches(1.0), pct, label, fill=(RED if is_last else SURFACE),
              line=RED if is_last else BORDER, label_color=(WHITE if is_last else INK), label_size=18)
    if i < len(fa_steps) - 1:
        add_arrow(s, x + fw, fy + Inches(0.5), x + fw + fgap, fy + Inches(0.5))

add_rect(s, Inches(0.8), Inches(5.15), Inches(11.7), Inches(1.55), fill=SURFACE, line=BORDER)
add_textbox(s, Inches(1.1), Inches(5.3), Inches(11), Inches(1.3),
            "Recommended production config: whofi_transformer + calibA + ~30s decision window.\n"
            "Raw's peak AUROC is higher (0.899 vs 0.821), but that's the wrong number to optimize for this "
            "task — calibA is dramatically better at the decision that actually matters: correctly "
            "rejecting a real intruder.",
            size=13.5, color=INK_SECONDARY)

add_footer(s, 11)

# =================================================================================================
# Slide 12 -- Known limitations
# =================================================================================================
s = add_slide()
add_kicker_title(s, "Known limitations", "Read this before demoing it", accent=RED)

add_bullets(s, Inches(0.8), Inches(1.4), Inches(5.7), Inches(5.2), [
    "taskB (identity, anjali/barath) tested BELOW CHANCE cross-day — AUROC 0.20–0.31, where 0.5 is a "
    "coin flip. Included because it was asked for; live_infer.py labels its output experimental/unreliable",
    "Day1 (186-sub, 40 MHz) vs Day2/3 (128-sub, 20 MHz) were captured on non-overlapping RF frequency "
    "bands — resampling + calibA compensate, but router channel/bandwidth pinning before future "
    "collection days is still the field's actual recommended fix, not a purely algorithmic one",
    "calibA requires a fresh per-day (or per-session, live) empty-room baseline — a brand-new "
    "environment with no quiet-room recording has no calibration to apply",
    "Deployable checkpoints have zero internal held-out split by design — their printed accuracy is a "
    "sanity check, not a generalization estimate; only the separate sweep runs and Day3 live testing "
    "are the trustworthy numbers",
], size=14)

add_bullets(s, Inches(6.85), Inches(1.4), Inches(5.65), Inches(5.2), [
    "Explicitly ruled out or not fully wired in yet:",
    ("CSI-ratio method (cancel phase noise via 2 Rx antennas) — needs 2+ antennas; this ESP32-S3 setup "
     "is 1x1, not applicable", 1),
    ("ℓ₁ per-packet gain normalization — implemented, unit-tested, not yet run through the full sweep "
     "as a real comparison arm", 1),
    ("RSSI fusion — implemented as a non-invasive wrapper, verified on a small smoke sample only, not "
     "a full sweep arm", 1),
    ("Doppler/velocity-domain features — the most cross-day-robust representation found so far in "
     "exploratory work (AUROC never flips below chance in either direction), but needs a much longer "
     "10–18s decision horizon and isn't part of the shipped pipeline", 1),
], size=14)

add_footer(s, 12)

# =================================================================================================
# Slide 13 -- Where this sits relative to later work
# =================================================================================================
s = add_slide()
add_kicker_title(s, "For context", "How this relates to the other (uncommitted) approaches", accent=BLUE)

rows = [
    ("", "Main branch (this deck)", "Home Model (RF)", "Signature + EVM"),
    ("Committed to main?", "Yes", "No — local/uncommitted", "No — local/uncommitted"),
    ("Model", "WhoFiTransformer (1-layer)", "Random Forest, 300 trees", "Transformer + Weibull/EVM"),
    ("Tasks", "Presence, Auth-gate, Motion, Identity (unreliable)", "Closed 2-class identity only", "Open-set accept/reject + ID"),
    ("Calibration", "calibA (empty-room z-score)", "None needed (RF invariant)", "None (found harmful here)"),
    ("Status", "Deployed / validated", "Experimental", "Experimental"),
]
table_x, table_y = Inches(0.8), Inches(1.35)
table_w, table_h = Inches(11.7), Inches(5.5)
gtable = s.shapes.add_table(len(rows), 4, table_x, table_y, table_w, table_h).table
gtable.columns[0].width = Inches(2.3)
gtable.columns[1].width = Inches(3.4)
gtable.columns[2].width = Inches(3.0)
gtable.columns[3].width = Inches(3.0)
for r, row in enumerate(rows):
    for c, val in enumerate(row):
        cell = gtable.cell(r, c)
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        cell.margin_left = Pt(8)
        cell.margin_right = Pt(8)
        cell.margin_top = Pt(4)
        cell.margin_bottom = Pt(4)
        para = cell.text_frame.paragraphs[0]
        run = para.add_run()
        run.text = val
        run.font.size = Pt(13 if r > 0 else 14)
        run.font.name = "Segoe UI"
        if r == 0:
            run.font.bold = True
            run.font.color.rgb = WHITE
            cell.fill.solid()
            cell.fill.fore_color.rgb = NAVY
        else:
            run.font.bold = (c == 0) or (c == 1)
            run.font.color.rgb = BLUE if c == 1 else INK_SECONDARY
            cell.fill.solid()
            cell.fill.fore_color.rgb = SURFACE if r % 2 else PAGE

add_footer(s, 13)

prs.save("ml/reports/wifi_csi_main_branch_model.pptx")
print("Saved ml/reports/wifi_csi_main_branch_model.pptx")
